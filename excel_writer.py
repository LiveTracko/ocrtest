"""Excel + JSON report generation with pandas + openpyxl.

All file writes are atomic (write to tmp file in the same dir, then
os.replace) so stopping the job mid-save never corrupts voters.xlsx —
the previous good Excel always survives.
"""
from __future__ import annotations

import json
import os
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _atomic_replace(tmp_path: Path, final_path: Path) -> None:
    os.replace(tmp_path, final_path)


def _tmp_in_same_dir(final_path: Path, suffix: str) -> Path:
    fd, tmp_name = tempfile.mkstemp(
        dir=str(final_path.parent), prefix=final_path.name + ".", suffix=suffix)
    os.close(fd)
    return Path(tmp_name)

VOTER_COLUMNS = [
    "Page Number", "Serial Number", "EPIC Number", "Name",
    "Relation Type", "Relation Name", "House Number",
    "Age", "Gender", "Confidence", "Validation Status", "Review Reason",
]

REVIEW_COLUMNS = ["Page", "Serial", "EPIC", "Name", "Issue", "Status"]


def _style_sheet(ws, text_cols: set[int] | None = None):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    for col in ws.columns:
        idx = col[0].column
        letter = get_column_letter(idx)
        maxlen = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[letter].width = min(max(14, maxlen + 3), 42)
    if text_cols:
        for r in range(2, ws.max_row + 1):
            for c in text_cols:
                ws.cell(row=r, column=c).number_format = "@"  # text: keeps leading zeros


def _salvaged_rows(payload: dict) -> list[dict]:
    """Fallback: pull raw AI voters when validated `voters` is empty.

    If schema validation rejected the whole page, process_page stores the
    raw Gemini output under `raw_gemini`. We must still show those rows in
    Excel (flagged REVIEW) instead of silently dropping AI data.
    """
    voters = payload.get("voters") or []
    if voters:
        return voters
    raw = payload.get("raw_gemini") or {}
    raw_voters = raw.get("voters") if isinstance(raw, dict) else None
    if not isinstance(raw_voters, list):
        return []
    out: list[dict] = []
    for rv in raw_voters:
        if not isinstance(rv, dict):
            continue
        row = dict(rv)
        row.setdefault("needs_review", True)
        if not row.get("review_reason"):
            row["review_reason"] = (
                (payload.get("error") or "; ".join(payload.get("issues", [])) or
                 "schema validation failed — raw AI row kept for review")[:500])
        # normalise confidence to allowed vocabulary so Excel stays clean
        if row.get("confidence") not in ("HIGH", "MEDIUM", "LOW"):
            row["confidence"] = "LOW"
        out.append(row)
    return out


def build_dataframes(page_payloads: list[dict]):
    voter_rows, review_rows = [], []
    for payload in page_payloads:
        page = payload.get("page_number")
        status = payload.get("status", "")
        rows = payload.get("voters") or _salvaged_rows(payload)
        salvaged = bool(rows) and not payload.get("voters")
        for v in rows:
            needs = bool(v.get("needs_review"))
            voter_rows.append({
                "Page Number": page,
                "Serial Number": v.get("serial_number"),
                "EPIC Number": str(v.get("epic_number")) if v.get("epic_number") is not None else None,
                "Name": v.get("name"),
                "Relation Type": v.get("relation_type"),
                "Relation Name": v.get("relation_name"),
                "House Number": str(v.get("house_number")) if v.get("house_number") is not None else None,
                "Age": v.get("age"),
                "Gender": v.get("gender"),
                "Confidence": v.get("confidence"),
                "Validation Status": "REVIEW" if needs else ("FAILED" if status == "FAILED" else "OK"),
                "Review Reason": v.get("review_reason"),
            })
            if needs:
                review_rows.append({
                    "Page": page,
                    "Serial": v.get("serial_number"),
                    "EPIC": v.get("epic_number"),
                    "Name": v.get("name"),
                    "Issue": v.get("review_reason") or "; ".join(payload.get("issues", [])[:2]) or "needs review",
                    "Status": "PENDING",
                })
        if status == "FAILED" and not rows:
            review_rows.append({"Page": page, "Serial": None, "EPIC": None,
                                "Name": None,
                                "Issue": payload.get("error") or "; ".join(payload.get("issues", [])) or "page failed",
                                "Status": "FAILED"})
        elif salvaged:
            # Page failed validation but AI rows were kept above — make the
            # cause visible in review.xlsx too (one extra line, not per-row).
            review_rows.append({"Page": page, "Serial": None, "EPIC": None,
                                "Name": None,
                                "Issue": (payload.get("error") or "; ".join(
                                    payload.get("issues", [])) or
                                    "validation failed — raw AI rows kept")[:500],
                                "Status": "RAW_KEPT"})
    voters_df = pd.DataFrame(voter_rows, columns=VOTER_COLUMNS)
    review_df = pd.DataFrame(review_rows, columns=REVIEW_COLUMNS)
    return voters_df, review_df


def _write_df_atomic(df, final_path: Path, sheet_name: str,
                     text_cols: set[int] | None = None) -> None:
    """Write one DataFrame to Excel atomically (tmp + replace)."""
    final_path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = _tmp_in_same_dir(final_path, suffix=".tmp.xlsx")
    try:
        with pd.ExcelWriter(tmp_path, engine="openpyxl") as w:
            df.to_excel(w, sheet_name=sheet_name, index=False)
            _style_sheet(w.sheets[sheet_name], text_cols=text_cols)
        _atomic_replace(tmp_path, final_path)
    finally:
        try:
            if tmp_path.exists():
                tmp_path.unlink()
        except OSError:
            pass


def write_excels(output_dir: Path, page_payloads: list[dict]) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    voters_df, review_df = build_dataframes(page_payloads)

    voters_path = output_dir / "voters.xlsx"
    review_path = output_dir / "review.xlsx"

    _write_df_atomic(voters_df, voters_path, "VOTERS", text_cols={3, 7})  # EPIC + House as text

    if len(review_df) == 0:
        review_df = pd.DataFrame(
            [{"Page": None, "Serial": None, "EPIC": None,
              "Name": None, "Issue": "No records need review", "Status": "OK"}])
    _write_df_atomic(review_df, review_path, "REVIEW", text_cols={3})

    return voters_path, review_path


def write_report(output_dir: Path, stats: dict) -> Path:
    from checkpoint import _atomic_write_text  # local import: no cycle at runtime
    output_dir.mkdir(parents=True, exist_ok=True)
    p = output_dir / "processing_report.json"
    _atomic_write_text(p, json.dumps(stats, indent=2))
    return p
